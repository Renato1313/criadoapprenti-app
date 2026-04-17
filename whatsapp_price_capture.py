#!/usr/bin/env python3
"""Script simples de WhatsApp com IA local (Ollama).

Fluxo:
1) Abre WhatsApp Web e conversa por numero.
2) Envia saudacao + solicitacao (gerada por IA ou fallback fixo).
3) Aguarda resposta.
4) Extrai S500/S10 por regex e usa IA como fallback.
5) Salva no Excel.
"""

from __future__ import annotations

import json
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Optional

import requests
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

# ====== CONFIGURACAO ======
NUMERO = "5521996697733"  # sem +, sem espaco, sem traco
EMPRESA = "Empresa Exemplo"

SAUDACAO = "Bom dia"
SOLICITACAO_FALLBACK = "16"

USAR_IA = True
OLLAMA_MODEL = "qwen2.5:3b"
OLLAMA_ENDPOINT = "http://localhost:11434/api/generate"
OLLAMA_TIMEOUT_SEGUNDOS = 60

ARQUIVO_XLSX = Path("saida/precos_combustivel.xlsx")
TIMEOUT_RESPOSTA_SEGUNDOS = 7200  # 2h
POLL_SEGUNDOS = 2
# ==========================


def to_float_br(raw: str) -> float:
    value = str(raw).strip()
    if "," in value and "." in value:
        value = value.replace(".", "").replace(",", ".")
    elif "," in value:
        value = value.replace(",", ".")
    return float(value)


def call_ollama_json(prompt: str) -> dict:
    response = requests.post(
        OLLAMA_ENDPOINT,
        json={
            "model": OLLAMA_MODEL,
            "prompt": prompt,
            "format": "json",
            "stream": False,
        },
        timeout=OLLAMA_TIMEOUT_SEGUNDOS,
    )
    response.raise_for_status()
    payload = response.json()
    return json.loads(payload["response"])


def gerar_solicitacao_ia(empresa: str, fallback: str) -> str:
    if not USAR_IA:
        return fallback

    prompt = (
        "Crie uma unica mensagem curta para WhatsApp em portugues.\n"
        "Objetivo: pedir preco atual de combustivel S500 e S10.\n"
        "Seja educado e objetivo, sem emoji.\n"
        f"Empresa: {empresa}\n"
        "Retorne SOMENTE JSON com chave 'mensagem'."
    )

    try:
        data = call_ollama_json(prompt)
        msg = str(data.get("mensagem", "")).strip()
        if msg:
            return msg
    except Exception as exc:  # noqa: BLE001
        print(f"[aviso] IA nao gerou solicitacao, usando fallback. Detalhe: {exc}")

    return fallback


def extract_prices_regex(texto: str) -> dict:
    s500 = re.search(r"(?i)\bS\s*500\b[^0-9]{0,20}([0-9]+(?:[.,][0-9]+)?)", texto)
    s10 = re.search(r"(?i)\bS\s*10\b[^0-9]{0,20}([0-9]+(?:[.,][0-9]+)?)", texto)

    out = {}
    if s500:
        out["S500"] = to_float_br(s500.group(1))
    if s10:
        out["S10"] = to_float_br(s10.group(1))
    return out


def extract_prices_ia(texto: str) -> dict:
    prompt = (
        "Extraia os valores de combustivel da mensagem abaixo.\n"
        "Retorne SOMENTE JSON no formato: "
        '{"S500": numero_ou_null, "S10": numero_ou_null}\n'
        "Mensagem:\n"
        f"{texto}"
    )
    data = call_ollama_json(prompt)

    result = {"S500": None, "S10": None}
    for key in ("S500", "S10"):
        if key in data and data[key] is not None:
            result[key] = to_float_br(str(data[key]))
    return result


def extract_prices(texto: str) -> tuple[Optional[float], Optional[float], str]:
    base = extract_prices_regex(texto)
    s500 = base.get("S500")
    s10 = base.get("S10")
    metodo = "regex"

    if s500 is not None and s10 is not None:
        return s500, s10, metodo

    if not USAR_IA:
        return s500, s10, metodo

    try:
        ai_out = extract_prices_ia(texto)
        if s500 is None:
            s500 = ai_out.get("S500")
        if s10 is None:
            s10 = ai_out.get("S10")
        if s500 is not None or s10 is not None:
            metodo = "regex+ia"
    except Exception as exc:  # noqa: BLE001
        print(f"[aviso] IA falhou na extracao, mantendo regex. Detalhe: {exc}")

    return s500, s10, metodo


def get_last_inbound_message(page):
    msgs = page.locator("div.message-in")
    count = msgs.count()
    if count == 0:
        return None

    last = msgs.nth(count - 1)
    payload = last.evaluate(
        """(el) => {
            const clean = (s) => (s || "").replace(/\\u200e/g, "").trim();

            const copyable = el.querySelector("div.copyable-text");
            const stamp =
              el.getAttribute("data-id") ||
              (copyable ? copyable.getAttribute("data-pre-plain-text") : "") ||
              "";

            const nodes = el.querySelectorAll(
              "div.selectable-text.copyable-text, div.copyable-text, span[data-lexical-text='true'], span[dir='auto']"
            );

            const parts = [];
            nodes.forEach((n) => {
              const t = clean(n.textContent);
              if (t) parts.push(t);
            });

            let text = "";
            if (parts.length) {
              text = clean(parts.join("\\n"));
            } else {
              text = clean(el.innerText);
            }

            return { stamp, text };
        }"""
    )

    text = (payload.get("text") or "").strip()
    if not text:
        return None

    key = f"{payload.get('stamp', '')}|{text}"
    return {"key": key, "text": text}


def wait_new_inbound(page, baseline_key, timeout_seconds=7200, poll_seconds=2):
    limite = time.time() + timeout_seconds
    while time.time() < limite:
        atual = get_last_inbound_message(page)
        if atual and atual["key"] != baseline_key:
            return atual
        page.wait_for_timeout(int(poll_seconds * 1000))
    raise TimeoutError("Nenhuma nova mensagem recebida dentro do tempo limite.")


def salvar_excel(
    empresa: str,
    numero: str,
    solicitacao_enviada: str,
    mensagem_recebida: str,
    s500: Optional[float],
    s10: Optional[float],
    metodo_extracao: str,
) -> None:
    ARQUIVO_XLSX.parent.mkdir(parents=True, exist_ok=True)

    if ARQUIVO_XLSX.exists():
        wb = load_workbook(ARQUIVO_XLSX)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Precos"
        ws.append(
            [
                "timestamp",
                "empresa",
                "telefone",
                "solicitacao_enviada",
                "mensagem_recebida",
                "S500",
                "S10",
                "metodo_extracao",
            ]
        )

    ws.append(
        [
            datetime.now().isoformat(timespec="seconds"),
            empresa,
            numero,
            solicitacao_enviada,
            mensagem_recebida,
            s500,
            s10,
            metodo_extracao,
        ]
    )
    wb.save(ARQUIVO_XLSX)


with sync_playwright() as p:
    context = p.chromium.launch_persistent_context(
        user_data_dir="perfil_whatsapp",
        headless=False,
    )

    page = context.pages[0] if context.pages else context.new_page()

    try:
        print("Abrindo conversa pelo numero...")
        page.goto(
            f"https://web.whatsapp.com/send?phone={NUMERO}&app_absent=0",
            wait_until="domcontentloaded",
            timeout=120000,
        )

        page.wait_for_timeout(12000)
        input("Se a conversa abriu, aperte ENTER aqui no prompt para continuar...")

        caixa = page.locator('footer div[contenteditable="true"]').first
        caixa.wait_for(timeout=20000)

        baseline = get_last_inbound_message(page)
        baseline_key = baseline["key"] if baseline else None

        solicitacao = gerar_solicitacao_ia(empresa=EMPRESA, fallback=SOLICITACAO_FALLBACK)
        print(f"Solicitacao escolhida: {solicitacao}")

        caixa.click()
        caixa.fill(SAUDACAO)
        caixa.press("Enter")

        page.wait_for_timeout(1500)
        caixa.fill(solicitacao)
        caixa.press("Enter")

        print("Mensagens enviadas com sucesso! Aguardando resposta...")

        resposta = wait_new_inbound(
            page,
            baseline_key=baseline_key,
            timeout_seconds=TIMEOUT_RESPOSTA_SEGUNDOS,
            poll_seconds=POLL_SEGUNDOS,
        )

        print("\nMensagem recebida:")
        print("-" * 40)
        print(resposta["text"])
        print("-" * 40)

        s500, s10, metodo = extract_prices(resposta["text"])
        print(f"Extraido -> S500: {s500} | S10: {s10} | metodo: {metodo}")

        salvar_excel(
            empresa=EMPRESA,
            numero=NUMERO,
            solicitacao_enviada=solicitacao,
            mensagem_recebida=resposta["text"],
            s500=s500,
            s10=s10,
            metodo_extracao=metodo,
        )
        print(f"Salvo em: {ARQUIVO_XLSX.resolve()}")
    except Exception as e:  # noqa: BLE001
        print(f"Erro: {e}")
    finally:
        input("Pressione ENTER para fechar...")
        context.close()
